import io, json, os, random, secrets, socket, time, uuid
from datetime import datetime
from functools import wraps

from dotenv import load_dotenv
from flask import Flask, g, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook, load_workbook
from sqlalchemy import (BigInteger, Boolean, Column, Float, ForeignKey, Integer, String, Unicode,
                        UnicodeText, create_engine, func, select)
from sqlalchemy.orm import declarative_base, scoped_session, sessionmaker
from werkzeug.security import check_password_hash, generate_password_hash

BASE = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE, 'cau-hinh.env'))
# Đổi CSDL chỉ cần sửa DATABASE_URL trong cau-hinh.env (mặc định: SQLite)
engine = create_engine(os.getenv('DATABASE_URL') or f"sqlite:///{os.path.join(BASE, 'data.db')}", pool_pre_ping=True)
S = scoped_session(sessionmaker(bind=engine))
Base = declarative_base()

# ---- Mô hình dữ liệu (Unicode/UnicodeText -> NVARCHAR trên SQL Server, lưu tiếng Việt đúng) ----
class User(Base):
    __tablename__ = 'users'
    id = Column(Integer, primary_key=True)
    username = Column(Unicode(100), unique=True, nullable=False)
    pass_hash = Column(String(255), nullable=False)
    full_name = Column(Unicode(200)); unit = Column(Unicode(200))
    role = Column(String(20), default='user')  # admin | user

class Sess(Base):
    __tablename__ = 'sessions'
    token = Column(String(64), primary_key=True)
    user_id = Column(Integer, ForeignKey('users.id'), nullable=False)

class Topic(Base):
    __tablename__ = 'topics'
    id = Column(Integer, primary_key=True)
    name = Column(Unicode(200), unique=True, nullable=False)

class Question(Base):
    __tablename__ = 'questions'
    id = Column(Integer, primary_key=True)
    topic_id = Column(Integer, ForeignKey('topics.id'))
    content = Column(UnicodeText, nullable=False)
    options = Column(UnicodeText, nullable=False)       # JSON: ["A...", "B...", ...]
    correct = Column(Integer, nullable=False)           # chỉ số đáp án đúng (0 = A)
    difficulty = Column(Integer, default=2)             # 1 dễ, 2 trung bình, 3 khó
    explanation = Column(UnicodeText)
    active = Column(Boolean, default=True)
    batch = Column(String(40))                          # lần import

class Exam(Base):
    __tablename__ = 'exams'
    id = Column(Integer, primary_key=True)
    title = Column(Unicode(300), nullable=False)
    kind = Column(String(20), default='mock')           # mock | official
    code = Column(Unicode(50))
    topic_ids = Column(String(500))                     # JSON list, rỗng = mọi chủ đề
    difficulty = Column(Integer, default=0)             # 0 = tất cả
    question_count = Column(Integer, nullable=False)
    duration_min = Column(Integer, default=0)           # 0 = vô hạn
    start_at = Column(String(32)); end_at = Column(String(32))
    status = Column(String(10), default='open')

class Attempt(Base):
    __tablename__ = 'attempts'
    id = Column(String(36), primary_key=True)
    exam_id = Column(Integer, ForeignKey('exams.id')); user_id = Column(Integer, ForeignKey('users.id'))
    mode = Column(String(20), nullable=False)           # practice | mock | official
    duration_min = Column(Integer, default=0)
    started_at = Column(BigInteger); submitted_at = Column(BigInteger)
    time_spent = Column(Integer); correct_count = Column(Integer); total = Column(Integer); score = Column(Float)

class Answer(Base):
    __tablename__ = 'attempt_answers'
    attempt_id = Column(String(36), ForeignKey('attempts.id'), primary_key=True)
    question_id = Column(Integer, ForeignKey('questions.id'), primary_key=True)
    pos = Column(Integer); chosen = Column(Integer); is_correct = Column(Boolean)

# ---- Tiện ích ----
app = Flask(__name__, static_folder=os.path.join(BASE, 'public'), static_url_path='')

@app.teardown_appcontext
def _close(_): S.remove()

def err(m, c=400): return jsonify(error=m), c
def I(x, d=0):
    try: return int(x)
    except (TypeError, ValueError): return d

@app.before_request
def _auth():
    g.user = None
    t = request.headers.get('Authorization', '')[7:]
    s = S.get(Sess, t) if t else None
    if s: g.user = S.get(User, s.user_id)

def admin_only(f):
    @wraps(f)
    def w(*a, **k):
        if not g.user: return err('Vui lòng đăng nhập', 401)
        if g.user.role != 'admin': return err('Không có quyền', 403)
        return f(*a, **k)
    return w

def new_session(u):
    t = secrets.token_hex(24); S.add(Sess(token=t, user_id=u.id)); S.commit(); return t

@app.get('/')
def index(): return send_from_directory(app.static_folder, 'index.html')

# ---- Tài khoản ----
@app.post('/api/register')
def register():
    b = request.get_json(silent=True) or {}
    un, pw, fn = (b.get('username') or '').strip(), b.get('password') or '', (b.get('full_name') or '').strip()
    if not un or not fn or len(pw) < 4: return err('Nhập đủ họ tên, tài khoản, mật khẩu (từ 4 ký tự)')
    if S.scalar(select(User).where(User.username == un)): return err('Tài khoản đã tồn tại')
    u = User(username=un, pass_hash=generate_password_hash(pw), full_name=fn, unit=b.get('unit') or '')
    S.add(u); S.commit()
    return jsonify(token=new_session(u))

@app.post('/api/login')
def login():
    b = request.get_json(silent=True) or {}
    u = S.scalar(select(User).where(User.username == (b.get('username') or '').strip()))
    if not u or not check_password_hash(u.pass_hash, b.get('password') or ''): return err('Sai tài khoản hoặc mật khẩu')
    return jsonify(token=new_session(u))

@app.get('/api/me')
def me():
    if not g.user: return err('Chưa đăng nhập', 401)
    return jsonify(id=g.user.id, username=g.user.username, full_name=g.user.full_name, role=g.user.role)

# ---- Lõi: rút đề, tạo & nạp bài làm ----
def pick(topic_ids, diff, count):
    q = select(Question.id).where(Question.active == True)  # noqa: E712
    if topic_ids: q = q.where(Question.topic_id.in_([I(x) for x in topic_ids]))
    if I(diff): q = q.where(Question.difficulty == I(diff))
    ids = list(S.scalars(q)); random.shuffle(ids)  # rút ngẫu nhiên ở Python: chạy giống nhau trên mọi CSDL
    return ids[:max(1, I(count, 10))]

def create_attempt(ids, mode, duration, exam_id=None, user_id=None):
    a = Attempt(id=str(uuid.uuid4()), exam_id=exam_id, user_id=user_id, mode=mode, duration_min=duration,
                started_at=int(time.time() * 1000), total=len(ids))
    S.add(a); S.flush()
    S.add_all(Answer(attempt_id=a.id, question_id=q, pos=i) for i, q in enumerate(ids))
    S.commit(); return a.id

def payload(aid):  # đề gửi thí sinh: KHÔNG kèm đáp án
    a = S.get(Attempt, aid)
    qs = S.scalars(select(Question).join(Answer, Answer.question_id == Question.id)
                   .where(Answer.attempt_id == aid).order_by(Answer.pos))
    left = max(0, a.duration_min * 60 - (int(time.time() * 1000) - a.started_at) // 1000) if a.duration_min else 0
    return {'id': aid, 'mode': a.mode, 'duration_min': a.duration_min, 'seconds_left': left,
            'questions': [{'id': q.id, 'content': q.content, 'options': json.loads(q.options)} for q in qs]}

@app.get('/api/topics')
def topics():
    r = S.execute(select(Topic.id, Topic.name, func.count(Question.id))
                  .outerjoin(Question, (Question.topic_id == Topic.id) & (Question.active == True))  # noqa: E712
                  .group_by(Topic.id, Topic.name).order_by(Topic.name))
    return jsonify([{'id': i, 'name': n, 'n': c} for i, n, c in r])

@app.post('/api/practice/start')  # luyện tập: không cần đăng nhập
def practice():
    b = request.get_json(silent=True) or {}
    ids = pick(b.get('topic_ids'), b.get('difficulty'), b.get('count'))
    if not ids: return err('Không có câu hỏi phù hợp với lựa chọn của bạn')
    d = I(b.get('duration_min'))
    return jsonify(payload(create_attempt(ids, 'practice', d if d in (30, 60) else 0)))

@app.get('/api/exams')
def exams():
    return jsonify([{'id': e.id, 'title': e.title, 'kind': e.kind, 'question_count': e.question_count,
                     'duration_min': e.duration_min, 'start_at': e.start_at, 'end_at': e.end_at, 'has_code': bool(e.code)}
                    for e in S.scalars(select(Exam).where(Exam.status == 'open').order_by(Exam.id.desc()))])

@app.post('/api/exams/<int:eid>/start')
def exam_start(eid):
    e, b, now = S.get(Exam, eid), request.get_json(silent=True) or {}, datetime.now()
    if not e or e.status != 'open': return err('Phòng thi không mở')
    try:
        if e.start_at and now < datetime.fromisoformat(e.start_at): return err('Chưa đến giờ mở phòng thi')
        if e.end_at and now > datetime.fromisoformat(e.end_at): return err('Phòng thi đã đóng')
    except ValueError: pass
    if e.code and e.code != b.get('code'): return err('Sai mã phòng thi')
    official = e.kind == 'official'
    if official and not g.user: return err('Thi chính thức yêu cầu đăng nhập', 401)
    if official:  # mỗi người thi 1 lần; làm dở thì vào lại đúng đề cũ
        old = S.scalar(select(Attempt).where(Attempt.exam_id == e.id, Attempt.user_id == g.user.id))
        if old: return err('Bạn đã hoàn thành bài thi này') if old.submitted_at else jsonify(payload(old.id))
    ids = pick(json.loads(e.topic_ids) if e.topic_ids else None, e.difficulty, e.question_count)
    if not ids: return err('Ngân hàng chưa có câu hỏi phù hợp')
    return jsonify(payload(create_attempt(ids, e.kind, e.duration_min, e.id, g.user.id if g.user else None)))

@app.post('/api/attempts/<aid>/submit')  # chấm điểm trên server
def submit(aid):
    a = S.get(Attempt, aid)
    if not a: return err('Không tìm thấy bài thi', 404)
    if a.user_id and (not g.user or g.user.id != a.user_id): return err('Không có quyền', 403)
    ans = (request.get_json(silent=True) or {}).get('answers') or {}
    rows = S.execute(select(Answer, Question).join(Question, Question.id == Answer.question_id)
                     .where(Answer.attempt_id == aid).order_by(Answer.pos)).all()
    if not a.submitted_at:
        c = 0
        for x, q in rows:
            v = ans.get(str(q.id))
            x.chosen = v if isinstance(v, int) and not isinstance(v, bool) else None
            x.is_correct = x.chosen == q.correct; c += int(x.is_correct)
        now = int(time.time() * 1000)
        a.submitted_at, a.time_spent, a.correct_count, a.total = now, round((now - a.started_at) / 1000), c, len(rows)
        a.score = round(c / len(rows) * 10, 1) if rows else 0
        S.commit()
    return jsonify(correct_count=a.correct_count, total=a.total, score=a.score, time_spent=a.time_spent,
                   review=[{'content': q.content, 'options': json.loads(q.options), 'correct': q.correct,
                            'chosen': x.chosen, 'explanation': q.explanation} for x, q in rows])

# ---- Quản trị ----
@app.get('/api/admin/template')
def template():
    wb = Workbook(); ws = wb.active; ws.title = 'CauHoi'
    ws.append(['Câu hỏi', 'A', 'B', 'C', 'D', 'E', 'Đáp án', 'Độ khó', 'Chủ đề', 'Giải thích'])
    ws.append(['Thủ đô của Việt Nam là thành phố nào?', 'Hà Nội', 'Huế', 'Đà Nẵng', 'TP. Hồ Chí Minh', '', 'A', 'Dễ',
               'Địa lý', 'Hà Nội là thủ đô của nước Cộng hòa xã hội chủ nghĩa Việt Nam.'])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='mau-cau-hoi.xlsx')

@app.post('/api/admin/import')
@admin_only
def import_questions():
    try:
        rows = list(load_workbook(io.BytesIO(request.get_data()), read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    except Exception: return err('Không đọc được file (chỉ nhận .xlsx)')
    if not rows: return err('File trống')
    head = [str(h or '').strip() for h in rows[0]]
    D = {'dễ': 1, 'de': 1, 'trung bình': 2, 'tb': 2, 'khó': 3, 'kho': 3}
    topics_by_name = {t.name: t for t in S.scalars(select(Topic))}
    batch, added, errors = datetime.now().isoformat(timespec='seconds'), 0, []
    for i, r in enumerate(rows[1:], start=2):
        d = {head[j]: ('' if v is None else str(v).strip()) for j, v in enumerate(r) if j < len(head)}
        if not any(d.values()): continue
        opts = [o for o in (d.get(k, '') for k in 'ABCDE') if o]
        a = d.get('Đáp án', '').upper()[:1]
        ans = 'ABCDE'.find(a) if a else -1
        if not d.get('Câu hỏi') or len(opts) < 2 or ans < 0 or ans >= len(opts):
            errors.append(f'Dòng {i}: thiếu câu hỏi/phương án hoặc đáp án không hợp lệ'); continue
        tn = d.get('Chủ đề') or 'Chung'
        if tn not in topics_by_name:
            t = Topic(name=tn); S.add(t); S.flush(); topics_by_name[tn] = t
        S.add(Question(topic_id=topics_by_name[tn].id, content=d['Câu hỏi'], options=json.dumps(opts, ensure_ascii=False),
                       correct=ans, difficulty=D.get(d.get('Độ khó', '').lower(), 2), explanation=d.get('Giải thích', ''), batch=batch))
        added += 1
    S.commit()
    return jsonify(added=added, errors=errors[:20])

@app.post('/api/admin/exams')
@admin_only
def exam_new():
    b = request.get_json(silent=True) or {}
    if not b.get('title') or I(b.get('question_count')) <= 0: return err('Cần nhập tên phòng và số câu hỏi')
    S.add(Exam(title=b['title'], kind='official' if b.get('kind') == 'official' else 'mock', code=b.get('code') or None,
               topic_ids=json.dumps([I(x) for x in b['topic_ids']]) if b.get('topic_ids') else None,
               difficulty=I(b.get('difficulty')), question_count=I(b['question_count']), duration_min=I(b.get('duration_min')),
               start_at=b.get('start_at') or None, end_at=b.get('end_at') or None))
    S.commit(); return jsonify(ok=1)

@app.get('/api/admin/exams')
@admin_only
def exam_list():
    out = []
    for e in S.scalars(select(Exam).order_by(Exam.id.desc())):
        done = S.scalar(select(func.count()).select_from(Attempt).where(Attempt.exam_id == e.id, Attempt.submitted_at.isnot(None)))
        out.append({'id': e.id, 'title': e.title, 'kind': e.kind, 'code': e.code, 'status': e.status, 'done': done})
    return jsonify(out)

@app.post('/api/admin/exams/<int:eid>/toggle')
@admin_only
def exam_toggle(eid):
    e = S.get(Exam, eid)
    if e: e.status = 'closed' if e.status == 'open' else 'open'; S.commit()
    return jsonify(ok=1)

@app.get('/api/admin/leaderboard/<int:eid>')  # bảng xếp hạng: chỉ admin
@admin_only
def leaderboard(eid):
    r = S.execute(select(User.full_name, User.username, User.unit, Attempt.score, Attempt.correct_count, Attempt.total, Attempt.time_spent)
                  .select_from(Attempt).outerjoin(User, User.id == Attempt.user_id)
                  .where(Attempt.exam_id == eid, Attempt.submitted_at.isnot(None))
                  .order_by(Attempt.score.desc(), Attempt.time_spent))
    return jsonify([dict(x._mapping) for x in r])

def init_db():
    try: Base.metadata.create_all(engine)
    except Exception as e:
        print('\nKHÔNG KẾT NỐI ĐƯỢC CƠ SỞ DỮ LIỆU:', str(e).strip().splitlines()[0][:300])
        print('Kiểm tra dòng DATABASE_URL trong cau-hinh.env (xem mục "Lỗi thường gặp" trong HUONG-DAN.md).')
        raise SystemExit(1)
    if not S.scalar(select(User).where(User.role == 'admin')):
        pw = os.getenv('ADMIN_PASSWORD') or 'admin123'
        S.add(User(username='admin', pass_hash=generate_password_hash(pw), full_name='Quản trị viên', role='admin')); S.commit()
        print(f'Tài khoản quản trị mặc định: admin / {pw}')

if __name__ == '__main__':
    from waitress import serve
    init_db(); port = I(os.getenv('PORT'), 3000)
    print(f'\nCơ sở dữ liệu: {engine.dialect.name}\nMáy này:  http://localhost:{port}')
    for ip in sorted({i for i in socket.gethostbyname_ex(socket.gethostname())[2] if not i.startswith('127.')}):
        print(f'Mạng LAN: http://{ip}:{port}')
    serve(app, host='0.0.0.0', port=port, threads=8)
